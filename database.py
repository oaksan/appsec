import os
from pathlib import Path
from typing import Generator
import openpyxl
from sqlalchemy import event
from sqlmodel import Session, SQLModel, create_engine, select

from models import (
    AssessmentAnswer,
    Squad,
    Statement,
    get_adoption_weight,
)

DB_FILE = "zeppelin.db"
SQLITE_URL = f"sqlite:///{DB_FILE}"

engine = create_engine(
    SQLITE_URL,
    connect_args={"check_same_thread": False},
)

@event.listens_for(engine, "connect")
def set_sqlite_pragma(dbapi_connection, connection_record):
    cursor = dbapi_connection.cursor()
    cursor.execute("PRAGMA journal_mode=WAL;")
    cursor.execute("PRAGMA foreign_keys=ON;")
    cursor.close()


def init_db(target_engine=None):
    """Inicializa as tabelas do banco de dados."""
    eng = target_engine or engine
    SQLModel.metadata.create_all(eng)


def get_session() -> Generator[Session, None, None]:
    """Dependency generator que provê uma sessão SQLModel por requisição."""
    with Session(engine) as session:
        yield session


def seed_database(session: Session, docs_dir: Path | str = "docs") -> None:
    """
    Executa a carga inicial automatizada e idempotente a partir dos arquivos
    docs/zeppelin_appsec_40_afirmacoes.xlsx e docs/simulacao_zeppelin_appsec_3_squads.xlsx.
    """
    # 1. Checagem de idempotência: se já temos 40 afirmações e 3 squads, não recarrega
    existing_statements = session.exec(select(Statement)).all()
    existing_squads = session.exec(select(Squad)).all()
    if len(existing_statements) == 40 and len(existing_squads) == 3:
        return

    base_path = Path(docs_dir)
    afirmacoes_path = base_path / "zeppelin_appsec_40_afirmacoes.xlsx"
    simulacao_path = base_path / "simulacao_zeppelin_appsec_3_squads.xlsx"

    if not afirmacoes_path.exists() or not simulacao_path.exists():
        raise FileNotFoundError(
            f"Arquivos de semente canônicos não encontrados em {base_path}. "
            f"Verifique se {afirmacoes_path.name} e {simulacao_path.name} estão presentes."
        )

    # 2. Carga das 40 Afirmações Atômicas
    wb_af = openpyxl.load_workbook(afirmacoes_path, data_only=True)
    ws_af = wb_af["Afirmações Atômicas"]

    statements_by_code: dict[str, Statement] = {}
    for r in range(2, 42):
        code = str(ws_af.cell(r, 1).value).strip()
        description = str(ws_af.cell(r, 2).value).strip()
        sth_stage = str(ws_af.cell(r, 3).value).strip()
        smaf_dimension = str(ws_af.cell(r, 4).value).strip()
        samm_ref = str(ws_af.cell(r, 5).value or "").strip()
        dsomm_ref = str(ws_af.cell(r, 6).value or "").strip()

        # Extrai código do estágio (A, B, C, D, E)
        stage_code = "A"
        if "Estágio " in sth_stage:
            stage_code = sth_stage.replace("Estágio ", "")[0]

        stmt = session.exec(select(Statement).where(Statement.code == code)).first()
        if not stmt:
            stmt = Statement(
                code=code,
                description=description,
                sth_stage=sth_stage,
                sth_stage_code=stage_code,
                smaf_dimension=smaf_dimension,
                samm_ref=samm_ref,
                dsomm_ref=dsomm_ref,
            )
            session.add(stmt)
            session.flush()
        statements_by_code[code] = stmt

    # 3. Definição das 3 Squads do Modelo Purdue
    squad_metadata = [
        {
            "code": "squad-a",
            "name": "Squad A",
            "sheet_name": "Squad A",
            "purdue_level": "Nível 2 — OT / Controle de Processos",
            "purdue_short": "N2 OT",
            "description": "Responsável por softwares de supervisão de campo (SCADA, IHMs) e controle em tempo real.",
        },
        {
            "code": "squad-b",
            "name": "Squad B",
            "sheet_name": "Squad B",
            "purdue_level": "Nível 3 — MES / Gestão de Operações de Manufatura",
            "purdue_short": "N3 MES",
            "description": "Sistemas de rastreabilidade de produção, controle de lotes e historiadores de dados, atuando na ponte entre o chão de fábrica (OT) e a TI corporativa.",
        },
        {
            "code": "squad-c",
            "name": "Squad C",
            "sheet_name": "Squad C",
            "purdue_level": "Nível 4 — ERP / TI Corporativa",
            "purdue_short": "N4 ERP",
            "description": "Gestão de sistemas de logística, portais corporativos e soluções em nuvem operando sob microsserviços e integração contínua.",
        },
    ]

    wb_sim = openpyxl.load_workbook(simulacao_path, data_only=True)

    for meta in squad_metadata:
        squad = session.exec(select(Squad).where(Squad.code == meta["code"])).first()
        if not squad:
            squad = Squad(
                code=meta["code"],
                name=meta["name"],
                purdue_level=meta["purdue_level"],
                purdue_short=meta["purdue_short"],
                description=meta["description"],
            )
            session.add(squad)
            session.flush()

        ws_squad = wb_sim[meta["sheet_name"]]
        # Linhas 4 a 43 contêm os 40 itens e suas avaliações
        for r in range(4, 44):
            item_code = str(ws_squad.cell(r, 1).value).strip()
            level_val = ws_squad.cell(r, 4).value
            try:
                adoption_level = int(level_val)
            except (ValueError, TypeError):
                adoption_level = 0

            adoption_weight = get_adoption_weight(adoption_level)
            stmt = statements_by_code.get(item_code)
            if not stmt:
                continue

            existing_ans = session.exec(
                select(AssessmentAnswer).where(
                    AssessmentAnswer.squad_id == squad.id,
                    AssessmentAnswer.statement_id == stmt.id,
                )
            ).first()

            if not existing_ans:
                ans = AssessmentAnswer(
                    squad_id=squad.id,
                    statement_id=stmt.id,
                    adoption_level=adoption_level,
                    adoption_weight=adoption_weight,
                )
                session.add(ans)

    session.commit()

